"""Compara duas planilhas .xlsx célula a célula (valores, mesclagens, painéis
congelados, listas suspensas, larguras). Uso: python3 comparar_planilhas.py a.xlsx b.xlsx"""
import sys
from openpyxl import load_workbook

def resumo(caminho):
    wb = load_workbook(caminho)
    out = {}
    for ws in wb.worksheets:
        vals = {}
        for row in ws.iter_rows():
            for c in row:
                if c.value not in (None, ""):
                    vals[c.coordinate] = c.value
        dv = sorted((str(d.sqref), d.formula1) for d in ws.data_validations.dataValidation)
        out[ws.title] = {
            "vals": vals, "merged": sorted(str(m) for m in ws.merged_cells.ranges), "freeze": ws.freeze_panes,
            "dv": dv, "filtro": ws.auto_filter.ref,
            "larg": {k: round(v.width, 1) for k, v in ws.column_dimensions.items() if v.width},
            "fills": {c.coordinate: c.fill.fgColor.rgb for row in ws.iter_rows(max_row=6) for c in row if c.fill and c.fill.fill_type == "solid"},
        }
    return wb.sheetnames, out

na, a = resumo(sys.argv[1])
nb, b = resumo(sys.argv[2])
print("abas iguais:", na == nb, na if na != nb else "")
dif = 0
for s in na:
    if s not in b:
        continue
    for k in ("vals", "merged", "freeze", "dv", "filtro", "larg", "fills"):
        x, y = a[s][k], b[s][k]
        if k == "vals":
            ks = set(x) | set(y)
            ruins = [kk for kk in sorted(ks) if x.get(kk) != y.get(kk)]
            if ruins:
                dif += 1
                print(f"[{s}] valores diferentes em {len(ruins)} células, ex.: " + "; ".join(f"{kk}: {x.get(kk)!r} x {y.get(kk)!r}" for kk in ruins[:3]))
        elif k == "fills":
            ruins = [kk for kk in set(x) | set(y) if (x.get(kk) or "")[-6:] != (y.get(kk) or "")[-6:]]
            if ruins:
                dif += 1
                print(f"[{s}] cores diferentes: " + "; ".join(f"{kk}: {x.get(kk)} x {y.get(kk)}" for kk in sorted(ruins)[:4]))
        elif x != y:
            dif += 1
            print(f"[{s}] {k}: {x} x {y}")
print("diferenças:", dif)
